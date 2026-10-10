#!/usr/bin/env python3
"""Parse uploaded .txt/.md/.html/.csv/.docx/.pdf into plain text."""
import sys, os, re, html as htmlmod

def parse_txt(path):
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            with open(path, 'r', encoding=enc) as f:
                return f.read()
        except Exception:
            continue
    return ""

def parse_html(path):
    raw = parse_txt(path)
    raw = re.sub(r'<(script|style)[^>]*>.*?</\1>', '\n', raw, flags=re.S | re.I)
    raw = re.sub(r'<br\s*/?>', '\n', raw, flags=re.I)
    raw = re.sub(r'</(p|div|h[1-6]|li|tr|ul|ol)>', '\n', raw, flags=re.I)
    raw = re.sub(r'<li[^>]*>', '- ', raw, flags=re.I)
    raw = re.sub(r'<h([1-6])[^>]*>', lambda m: '#' * int(m.group(1)) + ' ', raw, flags=re.I)
    raw = re.sub(r'<[^>]+>', '', raw)
    return htmlmod.unescape(raw)

def parse_csv(path):
    try:
        import csv
        out = []
        with open(path, 'r', encoding='utf-8-sig', newline='') as f:
            for row in csv.reader(f):
                out.append(' | '.join(c.strip() for c in row))
        return '\n'.join(out)
    except Exception:
        return parse_txt(path)

def parse_docx(path):
    try:
        from docx import Document
        doc = Document(path)
        parts = []
        for p in doc.paragraphs:
            t = (p.text or '').strip()
            if not t:
                continue
            style = (p.style.name or '').lower()
            if 'heading 1' in style:
                parts.append('# ' + t)
            elif 'heading 2' in style:
                parts.append('## ' + t)
            elif 'heading 3' in style:
                parts.append('### ' + t)
            elif 'list bullet' in style:
                parts.append('- ' + t)
            elif 'list number' in style:
                parts.append('1. ' + t)
            else:
                parts.append(t)
        for tbl in doc.tables:
            for row in tbl.rows:
                parts.append(' | '.join(c.text.strip() for c in row.cells))
        return '\n'.join(parts)
    except ImportError as e:
        print("need python-docx: " + str(e), file=sys.stderr)
        sys.exit(2)
    except Exception as e:
        print("docx parse failed: " + str(e), file=sys.stderr)
        sys.exit(1)

def parse_pdf(path):
    try:
        from pypdf import PdfReader
        r = PdfReader(path)
        return '\n'.join((pg.extract_text() or '') for pg in r.pages)
    except ImportError:
        try:
            from PyPDF2 import PdfReader
            r = PdfReader(path)
            return '\n'.join((pg.extract_text() or '') for pg in r.pages)
        except ImportError as e:
            print("need pypdf: " + str(e), file=sys.stderr)
            sys.exit(2)
    except Exception as e:
        print("pdf parse failed: " + str(e), file=sys.stderr)
        sys.exit(1)

if len(sys.argv) < 3:
    print("Usage: parse_upload.py <input_path> <ext>", file=sys.stderr)
    sys.exit(1)
inp, ext = sys.argv[1], sys.argv[2].lower().lstrip('.')
if ext == 'txt' or ext == 'md' or ext == 'markdown':
    print(parse_txt(inp))
elif ext == 'html' or ext == 'htm':
    print(parse_html(inp))
elif ext == 'csv':
    print(parse_csv(inp))
elif ext == 'docx':
    print(parse_docx(inp))
elif ext == 'pdf':
    print(parse_pdf(inp))
else:
    print("unsupported type: " + ext, file=sys.stderr)
    sys.exit(3)
